from pathlib import Path
from docx import Document

from generate_index_docs import configure, add_title, add_para, add_bullets, add_table, add_code


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "UEBA系统工作流与工作原理说明.docx"


def build():
    doc = Document()
    configure(doc, "UEBA 系统工作流与工作原理说明")
    add_title(doc, "UEBA 系统工作流与工作原理说明", "从异构日志接入到实体风险与调查闭环")
    add_para(doc, "本系统把 Windows、Zeek、Linux、VPN、DHCP、HR、AD 和 CMDB 等异构数据转换为可追溯的统一事件，再通过时态实体关系、行为特征、历史基线、异常检测和风险聚合形成可调查的 UEBA 结果。核心原则是事实、关系、特征、基线和风险分层保存；每层通过稳定 ID、时间范围和版本信息关联，任何异常都能回溯到原始证据。")

    doc.add_heading("1 系统目标", level=1)
    add_bullets(doc, [
        "把不同厂商对同一行为的描述转换为稳定、可跨来源消费的事件语义。",
        "识别人员、账号、设备、IP、应用等实体，并按事件发生时间建立关系。",
        "把离散事件转换为可建模的窗口特征，形成个人或同伴组行为基线。",
        "检测相对历史行为的偏离，把异常证据汇总为实体风险和调查案件。",
        "在字段缺失、时间不可信或关系不足时主动降级，避免输出误导性的无异常结论。",
    ])

    doc.add_heading("2 总体工作流", level=1)
    add_code(doc, "日志源与主数据\n  → 采集与原始证据归档\n  → Parser 解析和来源字段映射\n  → 统一事件入口 ingress\n  → 分类 classifier\n  → 质量校验 validate\n  → 领域路由 router\n  → 标准事件 Data Stream\n  → 实体解析与时态关系\n  → 窗口特征计算\n  → 行为基线训练与发布\n  → 异常检测和解释\n  → 风险聚合与当前风险\n  → 案件调查和反馈\n  → 映射 模型 规则持续迭代")
    stages = [
        ("1", "采集", "日志、快照和会话数据", "来源信封与原始证据", "保证不丢失、可重放"),
        ("2", "解析", "原始文本或 JSON", "来源字段与标准候选字段", "识别格式、类型和来源版本"),
        ("3", "分类校验路由", "标准化输入", "唯一领域事件或隔离事件", "建立语义、质量和物理位置"),
        ("4", "实体解析", "标准事件与主数据", "实体当前视图和时态关系", "回答事件发生时是谁、哪台设备"),
        ("5", "特征计算", "事件与关系", "实体窗口特征", "把事件转成可比较数值或集合"),
        ("6", "基线训练", "历史特征", "个人或同伴组基线", "学习正常行为范围"),
        ("7", "异常检测", "当前特征与基线", "异常、解释和证据引用", "识别显著偏离"),
        ("8", "风险聚合", "异常和上下文", "风险流水与当前风险", "跨时间、场景汇总风险"),
        ("9", "调查反馈", "异常、风险和原始证据", "案件、处置和标签", "闭环验证并改进规则模型"),
    ]
    add_table(doc, ["阶段", "处理", "输入", "输出", "主要作用"], stages, [0.45, 1.0, 1.55, 1.9, 1.7], 8.0)

    doc.add_heading("3 数据接入与原始证据", level=1)
    add_para(doc, "采集层接收行为日志和主数据。行为日志包括认证、网络、DNS、Web、进程、文件、会话和告警；HR、AD 与 CMDB 提供人员、账号、组织和资产的权威属性。主数据不参与普通行为计数，只用于实体当前视图、同伴组和关系解析。")
    add_table(doc, ["来源", "典型数据", "系统用途", "关键限制"], [
        ("Windows", "4624 4625 4634 4672 4688 账号组变更", "认证、会话、进程、IAM", "必须区分行为用户、操作者和目标账号"),
        ("Zeek", "conn dns http ssl files notice", "网络、DNS、Web、TLS、文件和告警", "本身通常不能证明用户身份"),
        ("VPN DHCP", "会话、分配 IP、租约", "用户 设备 IP 时态关系", "开始结束和租约有效期必须可靠"),
        ("HR AD", "人员、账号、部门、角色和状态", "账号到人员、同伴组、特权属性", "使用稳定 ID 和生效时间"),
        ("CMDB", "资产 ID、主机名、类型、关键性、共享标志", "设备实体和风险上下文", "不能只用当前 IP 代替设备身份"),
    ], [1.0, 2.0, 1.7, 1.9], 8.2)
    add_para(doc, "完整原文优先写入对象存储，也可短期写入 logs-ueba.raw-<namespace>。标准事件只保存 raw_event_id、raw_uri、raw_sha256 等引用，使后续任何结果可以回到不可变证据，同时避免每层复制整份日志。")

    doc.add_heading("4 解析 标准化和路由原理", level=1)
    add_para(doc, "所有生产者写入 logs-ueba.ingress-<namespace>。入口 Pipeline 串联分类、校验和路由。最终领域 Data Stream 不配置入口默认 Pipeline，因此 reroute 后不会再次进入解析链或形成循环。")
    add_code(doc, "vendor.dataset + event.code\n  否则 vendor.dataset + record_type 或 event_type\n  否则固定 vendor.dataset\n  否则 quarantine")
    add_table(doc, ["字段层", "代表字段", "作用"], [
        ("来源合同", "vendor.name product dataset schema_version payload", "标识输入格式并保存未晋升的长尾字段"),
        ("ECS 事实", "event user host source destination network process file", "表达跨来源一致的事实"),
        ("UEBA 语义", "ueba.event.type semantic_tags session time", "提供稳定事件类型、角色和时序语义"),
        ("受控路由", "ueba.route.domain rule_id version dataset", "记录由哪条规则进入哪个存储域"),
        ("质量", "ueba.quality.status usable_for errors warnings", "决定事件能支持哪些实体、特征和检测"),
        ("可追溯", "ueba.provenance.parser mapping raw 引用", "支持回放、审计和版本对比"),
    ], [1.05, 3.1, 2.45], 8.2)
    add_para(doc, "一个事件只选择一个主存储域。event.category、event.type 和 ueba.event.type 用于表达语义，不直接控制物理路由；跨领域用途通过 semantic_tags 和 usable_for 表达，避免同一事实被重复写入和重复计数。")

    doc.add_heading("5 标准事件存储", level=1)
    add_table(doc, ["事件域", "索引模式", "主要内容"], [
        ("认证", "logs-ueba.authentication-*", "登录、失败、票据和凭据验证"),
        ("身份权限", "logs-ueba.iam-* / directory-*", "账号、组、权限和目录对象操作"),
        ("终端", "logs-ueba.endpoint-* / file-*", "进程、服务、软件包和文件行为"),
        ("网络", "logs-ueba.network-* / dns-* / web-* / tls-*", "连接、域名、HTTP 和 TLS"),
        ("会话", "logs-ueba.session-*", "VPN、DHCP、堡垒机和地址分配"),
        ("告警状态指标", "logs-ueba.alert-* / state-* / metric-*", "外部告警、资产观察和采集健康"),
        ("隔离", "logs-ueba.quarantine-*", "无效、不支持或无法分类的输入"),
    ], [1.15, 2.75, 2.7], 8.2)
    add_para(doc, "统一事件模板由公共组件与一个领域组件组合。公共组件负责时间、租户、事件元数据、质量、证据和路由；领域组件只增加该领域需要的字段。Mapping 默认严格，未声明字段不能静默扩张索引。")

    doc.add_heading("6 数据质量和能力门禁", level=1)
    add_table(doc, ["状态", "含义", "下游处理"], [
        ("qualified", "主语义明确并满足事件合同", "允许 ready 能力消费"),
        ("partial", "存在非致命缺失、推断或冲突", "只允许明确支持的能力消费"),
        ("invalid", "时间、类型、格式或证据不可信", "禁止进入特征、基线和风险"),
        ("unsupported", "来源可识别但尚无发布映射", "保留证据并进入映射待办"),
    ], [1.05, 3.0, 2.55], 8.3)
    add_para(doc, "质量状态描述数据，不等于风险分数。检测任务必须同时检查窗口事件数、必需字段覆盖、事件时间质量、实体关系版本和 use case capability。条件不足时输出“数据不可用”或“降级”，不能输出“未发现异常”。")

    doc.add_heading("7 实体解析与时态关系", level=1)
    add_para(doc, "统一事件中的 user、host 和 IP 只是事件事实。实体解析服务把这些标识与 HR、AD、CMDB、VPN 和 DHCP 数据关联，生成稳定的 person、account、device、ip、application 实体，以及带有效期和置信度的关系。")
    add_code(doc, "account SID → account → person\nhostname 或 asset ID → device\nDHCP leased IP ↔ device [valid_from valid_to)\nVPN user ↔ assigned IP [session start session end)\nWindows logon → account logged_on device")
    add_table(doc, ["关系字段", "意义"], [
        ("relation.source / target", "关系两端的稳定 entity.id 和 entity.type"),
        ("relation.valid_from / valid_to", "关系在什么时间范围有效"),
        ("relation.confidence", "归属证据的可信程度"),
        ("relation.evidence_event_ids", "产生关系的 Windows、VPN、DHCP 或主数据证据"),
        ("relation.resolution_version", "实体解析规则版本，用于重算与审计"),
    ], [2.5, 4.1], 8.5)
    add_para(doc, "网络事件只有 IP 时，系统首先生成 IP 或设备级行为。只有事件发生时刻存在足够置信度的用户 设备 IP 关系，才允许把异常归属到人员。共享设备、NAT、跳板机和地址复用不能被强制映射到单一用户。")

    doc.add_heading("8 特征计算与行为基线", level=1)
    add_para(doc, "特征计算把事件流转换为实体在固定时间窗口内的数值、集合或类别统计。每条特征保存 feature.id、版本、实体、窗口、样本量、输入事件引用、Mapping 版本、关系解析版本和质量。")
    add_table(doc, ["特征类型", "示例", "工作原理"], [
        ("计数", "每小时成功登录数、失败数、进程启动数", "按实体与窗口聚合事件"),
        ("去重数", "来源 IP、设备、目的地、域名数量", "计算 cardinality 或确定性集合"),
        ("首次出现", "新设备、新目的地、新进程", "与历史观察集合比较"),
        ("时段分布", "登录小时、工作日和周末", "建立周期性行为直方图"),
        ("数值流量", "外发字节、上传正文长度", "按明确方向和协议层次求和"),
        ("关系特征", "同一 IP 涉及多个设备、关系置信度", "结合事件时态关系计算"),
    ], [1.2, 2.6, 2.8], 8.2)
    add_para(doc, "基线由历史特征训练，可针对单个实体或同伴组保存均值、标准差、分位数、样本量、训练区间和模型版本。训练窗口必须与当前评分窗口隔离，避免把待检测行为提前学习为正常。冷启动、样本不足和节假日等情况通过训练质量和场景策略处理。")

    doc.add_heading("9 异常检测和解释", level=1)
    add_para(doc, "检测器读取当前特征、已发布基线和场景合同，判断当前值是否显著偏离个人或同伴组历史。规则型检测也可以直接消费统一事件，例如敏感组成员变化。异常对象不仅保存分数，还保存当前值、期望值、偏离、原因码、规则模型版本及证据引用。")
    add_table(doc, ["场景", "主要输入", "异常依据"], [
        ("异常登录时间", "认证事件和时段特征", "登录时段偏离个人或同伴基线"),
        ("登录失败量", "失败认证窗口计数", "失败量、来源数或失败后成功异常"),
        ("非常用设备登录", "认证、设备实体和关系", "设备首次出现或历史频率过低"),
        ("新外部目的地", "网络连接和设备/IP 关系", "外部目的地首次出现或稀有"),
        ("异常外发量", "方向可靠的字节特征", "当前外发量显著超过历史范围"),
        ("疑似分片上传", "Web、Network、File 事件", "分片数量、字节、并发和新目的地异常；缺会话 ID 时低置信"),
        ("新进程", "终端进程事件", "进程或父子组合首次出现或稀有"),
        ("敏感组变化", "IAM 或目录变更", "敏感组、操作者、目标与审批上下文"),
    ], [1.55, 2.4, 2.65], 8.1)

    doc.add_heading("10 风险聚合与案件", level=1)
    add_para(doc, "异常不直接等同于威胁。风险引擎根据异常严重度、实体关键性、证据质量、时间衰减和多场景关联生成不可变 risk event，并更新 entity risk current。风险流水保留 score_delta、变化前后分数、原因码、策略版本和证据，当前风险索引提供快速查询视图。")
    add_code(doc, "罕见设备登录\n  + 大量读取敏感文件\n  + 向新外部目的地异常上传\n  → 多场景风险提升\n  → 创建或关联调查案件\n  → 分析员回查异常 特征 事件 原文")
    add_para(doc, "案件保存状态、严重度、负责人、相关实体、异常、风险事件、原始事件以及分析反馈。确认、误报、抑制和原因标签进入反馈闭环，用于调整映射、实体解析、特征、基线、规则和风险策略。")

    doc.add_heading("11 数据对象之间的关联", level=1)
    add_table(doc, ["上游对象", "关联键", "下游对象"], [
        ("原始证据", "raw_event_id raw_uri raw_sha256", "统一事件"),
        ("统一事件", "event.id", "关系、特征、异常"),
        ("实体", "entity.id", "关系、特征、异常、风险和案件"),
        ("关系", "relation.id resolution_version", "特征与异常"),
        ("特征", "feature.id version window.id", "基线和异常"),
        ("基线", "baseline.id model.version", "异常"),
        ("异常", "anomaly.id", "风险流水和案件"),
        ("风险流水", "risk.id calculation.version", "当前风险和案件"),
    ], [1.4, 2.65, 2.55], 8.3)
    add_para(doc, "派生对象只保存必要的输入引用和计算结果，不复制全部事件。Mapping 或关系修正后，系统根据版本和输入引用定位受影响窗口，写入新版本结果或影子索引，比较完成后再切换读取。")

    doc.add_heading("12 失败处理和重放", level=1)
    add_table(doc, ["失败类型", "落点", "恢复方式"], [
        ("业务语义 invalid 或 unsupported", "quarantine", "补充映射或修复数据后重新写入 ingress"),
        ("Mapping 类型等索引失败", "Elasticsearch Failure Store", "修复模板或 Pipeline 后重放"),
        ("partial", "正常领域 Data Stream", "限制 capability 并持续监控缺失率"),
        ("关系冲突", "关系质量告警", "人工或规则修正，定向重算相关窗口"),
        ("模型或规则升级", "新版本或影子输出", "历史回放、结果比较、Alias 切换"),
    ], [1.85, 1.9, 2.85], 8.2)
    add_para(doc, "重放必须保持 event.id 和派生对象幂等键稳定。同一输入重复执行不能增加独立动作数，也不能重复生成关系、特征、异常或风险。")

    doc.add_heading("13 运行监控", level=1)
    add_bullets(doc, [
        "按租户、来源、dataset、来源版本、Parser 和 Mapping 版本统计接收、归档、解析和路由数量。",
        "监控核心字段覆盖率、时间回退率、重复率、处理延迟和 Failure Store 增量。",
        "监控用户 设备 IP 关系覆盖率、关系冲突和人工抽样准确率。",
        "按场景展示 ready、degraded、unavailable 和 unknown，状态下降时停止或降级相关评分。",
        "检测任务输入为零不算健康，必须区分确实没有行为与数据链路不可用。",
    ])

    doc.add_heading("14 发布和版本管理", level=1)
    add_para(doc, "Parser、Mapping、Route Registry、统一 Schema、质量规则和测试集作为同一个发布包。制品使用不可变版本名；新增兼容字段通常升级 Minor 版本，字段类型、语义、数组形态或必填规则变化属于 Breaking 变化。")
    add_code(doc, "编译路由配置\n  → JSON 与模板静态校验\n  → Pipeline simulate\n  → 正常 异常 缺字段样例回放\n  → 影子运行与下游回归\n  → 切换写入或读取 Alias\n  → 观察指标\n  → 停用旧版本")

    doc.add_heading("15 端到端示例", level=1)
    add_para(doc, "某用户凌晨通过新的公网 IP 登录 VPN，获得一个内网地址，随后访问从未连接的外部目的地并产生异常外发流量。系统的处理过程如下。")
    add_table(doc, ["步骤", "处理结果"], [
        ("采集", "VPN、Windows、DHCP 和 Zeek 事件分别归档并获得稳定 event.id"),
        ("标准化", "登录进入 authentication，VPN/DHCP 进入 session，连接进入 network"),
        ("关系", "根据会话和租约在事件时刻关联 account person device IP"),
        ("特征", "计算非工作时段登录、新来源、新目的地和外发字节"),
        ("基线", "与用户或设备历史时段、来源、目的地和流量分布比较"),
        ("异常", "分别生成异常登录时间、新来源、新目的地和异常外发量证据"),
        ("风险", "同一实体短时间出现多项相关异常，风险分提高"),
        ("调查", "分析员从案件下钻到风险、异常、特征、关系、标准事件和原始日志"),
    ], [1.15, 5.45], 8.5)

    doc.add_heading("16 系统工作原理总结", level=1)
    add_bullets(doc, [
        "语义统一：用 ECS 事实字段和 ueba.event.type 消除厂商字段差异。",
        "时态归属：用有效期关系回答行为发生时账号、人员、设备和 IP 的真实对应。",
        "行为建模：把事件聚合为窗口特征，再与个人或同伴历史基线比较。",
        "证据驱动：异常和风险保存输入引用、计算值、版本与解释，而不是只给一个分数。",
        "质量门禁：数据不足时降级或停止检测，防止静默失败和错误归属。",
        "版本可重算：所有处理阶段记录版本，支持历史回放、影子比较和定向重算。",
    ])
    return doc


if __name__ == "__main__":
    document = build()
    document.save(OUT)
    print(OUT)
