from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

BLUE = "1F4E78"
PALE = "EAF2F8"
GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_keep_with_next(p, value=True):
    p.paragraph_format.keep_with_next = value


def set_cn_font(run, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def configure(doc, title):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, before, after in (("Title", 25, 0, 16), ("Heading 1", 17, 18, 8), ("Heading 2", 13, 12, 5), ("Heading 3", 11.5, 9, 4)):
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(size)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    styles["Title"].font.bold = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(title)
    set_cn_font(r)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 100, 100)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(subtitle)
    r.bold = True
    r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(13)


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
        set_cn_font(r)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
            if ridx % 2:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            set_cn_font(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(9)


def doc_structure():
    doc = Document()
    configure(doc, "UEBA 后端 Index 表结构设计")
    add_title(doc, "UEBA 后端 Index 表结构设计", "基于 elasticsearch-storage-full 1.1.0 的实现说明")
    add_para(doc, "本文说明当前后端如何用 Elasticsearch Data Stream、普通索引、组件模板和路由管道组织 UEBA 数据。结论是：统一事件按业务域分流，实体和当前状态使用可更新普通索引，关系、特征、异常和风险流水使用追加型 Data Stream；所有派生对象通过稳定 ID、时间窗口和版本字段关联，不复制整份原始日志。")

    doc.add_heading("1 设计目标和边界", level=1)
    add_bullets(doc, [
        "用稳定语义合同支持跨来源查询、实体解析、特征计算、异常检测和调查取证。",
        "按对象的写入模型和生命周期选存储形态，不按厂商、用户或单台设备建索引。",
        "对高价值字段建立显式 Mapping；长尾来源字段放入 vendor.payload flattened，原文独立保存。",
        "用 strict Mapping、字段数量上限、版本化模板和隔离流控制 Mapping Explosion 与不兼容变更。",
    ])

    doc.add_heading("2 索引总体分层", level=1)
    rows = [
        ("原始证据", "logs-ueba.raw-<namespace>", "Data Stream", "@timestamp + event.id", "回放、审计、取证", "30 天"),
        ("统一入口", "logs-ueba.ingress-<namespace>", "Data Stream", "@timestamp", "分类、校验、路由；不长期保存", "入口型"),
        ("统一事件", "logs-ueba.<domain>-<namespace>", "Data Stream", "@timestamp + event.id", "跨来源标准行为事实", "90 天"),
        ("隔离事件", "logs-ueba.quarantine-<namespace>", "Data Stream", "@timestamp", "invalid unsupported 和业务错误", "按治理策略"),
        ("实体当前视图", "ueba-entities-v<schema>-<namespace>", "普通索引 + Alias", "entity.id", "人员、账号、设备、IP、应用当前状态", "长期当前态"),
        ("时态关系", "ueba-entity-relations-<namespace>", "Data Stream", "@timestamp + relation.id", "用户 账号 设备 IP 的有效期关系", "365 天"),
        ("行为特征", "ueba-features-<namespace>", "Data Stream", "实体 + 窗口 + feature.version", "窗口统计和维度特征", "180 天"),
        ("行为基线", "ueba-baselines-v<schema>-<namespace>", "普通索引 + Alias", "baseline.id", "当前发布基线和训练元数据", "版本保留"),
        ("异常", "ueba-anomalies-<namespace>", "Data Stream", "@timestamp + anomaly.id", "异常结果、解释和证据", "365 天"),
        ("风险流水", "ueba-risk-events-<namespace>", "Data Stream", "@timestamp + risk.id", "风险增量、衰减、关联", "365 天"),
        ("当前风险", "ueba-entity-risk-current-v<schema>-<namespace>", "普通索引 + Alias", "entity.id", "实体最新风险快照", "长期当前态"),
        ("案件", "ueba-cases-v<schema>-<namespace>", "普通索引", "case.id", "调查状态、处置、反馈", "业务留存"),
    ]
    add_table(doc, ["对象", "索引模式", "形态", "主键或时间键", "职责", "默认周期"], rows, [0.72, 1.58, 0.78, 1.18, 1.75, 0.72], 7.5)

    doc.add_heading("3 统一事件索引结构", level=1)
    add_para(doc, "统一事件模板采用“公共组件 + 单一领域组件”的组合。公共组件管理所有事件都需要的时间、租户、来源、质量、证据和版本字段；领域组件只增加该领域可被查询、关联或检测使用的 ECS 字段。")
    add_code(doc, "logs-ueba.authentication-*\n  = ueba-normalized-event-common@1.1.0\n  + ueba-normalized-event-authentication@1.1.0")
    domains = [
        ("authentication", "用户、目标用户、组、主机、源/目的地址、client/server", "Windows 登录、VPN、IdP、Linux SSH"),
        ("iam", "用户、目标用户、组、主机、源地址", "账号、组、权限变更"),
        ("directory", "用户、组、主机、源地址、文件/目录对象", "AD LDAP 对象访问与变更"),
        ("endpoint", "主机、用户、进程、父进程、文件、软件包", "进程、服务、注册表、模块"),
        ("network", "源/目的端点、方向、协议、字节包数、观察者", "Zeek conn、防火墙、NDR"),
        ("dns", "网络端点、观察者、DNS 问题与响应", "查询名、类型、响应码、解析 IP"),
        ("web", "网络端点、HTTP、URL、User Agent、观察者", "HTTP 请求和响应"),
        ("tls", "网络端点、TLS 版本/套件/状态、观察者", "TLS 会话与证书观察"),
        ("file", "主机、用户、网络端点、文件与哈希", "文件访问、创建、传输"),
        ("session", "用户、主机、网络端点、UEBA 会话", "VPN、DHCP、堡垒机会话"),
        ("alert", "主机、用户、网络端点、规则", "Zeek notice、EDR、IDS、DLP 告警"),
        ("state", "主机、用户、服务、软件包、观察者", "资产和服务状态观察"),
        ("metric", "主机、服务、观察者", "采集、传感器和 Pipeline 指标"),
    ]
    add_table(doc, ["领域", "领域字段组", "典型事实"], domains, [1.05, 3.2, 2.35], 8.2)

    doc.add_heading("4 公共字段合同", level=1)
    common = [
        ("身份与时间", "@timestamp event.id event.created event.ingested event.start event.end event.duration", "窗口计算、去重、排序和回放"),
        ("事件语义", "event.kind category type action outcome code dataset；ueba.event.type semantic_tags", "跨来源分类与检测合同"),
        ("租户与位置", "organization.id；data_stream.type dataset namespace", "隔离、授权和物理落点"),
        ("来源合同", "vendor.name product dataset schema_version payload", "识别输入格式并保留未晋升字段"),
        ("路由", "ueba.route.domain dataset rule_id version", "证明由哪条受控规则落到哪个域"),
        ("质量", "ueba.quality.status score errors warnings usable_for", "按用例判断 qualified partial invalid unsupported"),
        ("可追溯性", "ueba.provenance.raw_event_id raw_uri raw_sha256 parser mapping 版本", "从派生事实回查原始证据和处理版本"),
        ("时间质量", "ueba.time.source quality original original_timezone", "识别准确时间或 fallback"),
        ("长尾与原文", "vendor.payload flattened；event.original index false", "防字段爆炸；保留回查信息"),
    ]
    add_table(doc, ["字段组", "代表字段", "设计用途"], common, [1.15, 3.65, 1.8], 8.1)

    doc.add_heading("5 派生索引结构", level=1)
    derived = [
        ("entities", "entity.id type subtype name status criticality attributes source_refs identifiers.* version resolution.*", "以 entity.id 做幂等 upsert；不同标识归并到稳定实体"),
        ("entity-relations", "relation.id type source target valid_from valid_to confidence evidence_event_ids resolution_version", "按事件时间解析 logged_on assigned_ip owns uses 等关系"),
        ("features", "entity.* feature.id/version value values unit dimensions sample_count window.* inputs.* quality.*", "一条记录表达某实体某时间窗的一项特征"),
        ("baselines", "baseline.* training.* model.* versions.*", "保存实体或同伴组对某特征的统计分布和模型版本"),
        ("anomalies", "anomaly.* entity.* feature.* baseline.* evidence.* explanation.* detection.* quality.*", "保存当前值、期望值、偏离、原因和证据"),
        ("risk-events", "risk.id type score_delta before after level reason_codes expires_at entity.* evidence.* calculation.*", "不可变风险账本，支持回放、衰减和审计"),
        ("entity-risk-current", "entity.* risk.score level status first_seen last_updated top_reason_codes latest.* calculation.*", "按实体覆盖更新的快速查询视图"),
        ("cases", "case.* feedback.* entities.* evidence.* version", "调查工作流与模型反馈闭环"),
    ]
    add_table(doc, ["索引", "核心字段", "写入语义"], derived, [1.22, 3.55, 1.83], 8.0)

    doc.add_heading("6 写入链路和关联规则", level=1)
    add_code(doc, "来源原文 → raw evidence\n标准化输入 → ingress → classifier → validate → router\n             → logs-ueba.<domain>-*\n             → entity resolution → entities / entity-relations\n             → feature jobs → features / baselines\n             → detection → anomalies → risk-events → entity-risk-current\n             → analyst workflow → cases")
    add_bullets(doc, [
        "生产者只写 ingress；分类器按 vendor.dataset 与 event.code 或 record_type 匹配注册表。",
        "event.category、event.type 和 ueba.event.type 只表达语义，不直接决定物理路由。",
        "派生记录必须保存输入 event.id、relation.id、feature.id 以及 mapping/resolution/model/rule 版本。",
        "同一事件只落一个主事件域；多用途通过 semantic_tags 和 usable_for 表达，避免重复计数。",
    ])

    doc.add_heading("7 Mapping 与容量约束", level=1)
    controls = [
        ("动态策略", "统一事件和派生对象采用 strict；quarantine 可容纳异常输入", "未知字段不能静默扩张生产 Mapping"),
        ("字段上限", "统一事件 2500；实体/关系/特征/基线/风险多为 400 到 600；原始证据 200", "限制 Mapping Explosion"),
        ("对象深度", "统一事件 depth limit 20，nested fields limit 20", "控制复杂对象结构"),
        ("文本类型", "精确键 keyword；长检索文本 match_only_text；命令行和 URL wildcard", "与查询方式匹配"),
        ("数值类型", "计数 long；测量值 double；置信度和质量分 scaled_float", "兼顾聚合精度与存储"),
        ("副本", "Data Stream 原型 auto_expand_replicas 0-1", "单节点可运行，多节点自动增加一个副本"),
    ]
    add_table(doc, ["控制项", "当前设计", "目的"], controls, [1.15, 3.6, 1.85], 8.2)

    doc.add_heading("8 生命周期与版本策略", level=1)
    add_bullets(doc, [
        "raw evidence 默认 30 天；normalized events 默认 90 天；features 默认 180 天；relations 和 analysis records 默认 365 天。实际值需按合规、训练窗口与调查回溯期校准。",
        "可更新对象通过带 schema 版本的物理索引和稳定 Alias 发布；Data Stream 通过不可变组件模板版本演进。",
        "新增可选字段属于 Minor 版本；改类型、改语义、删除字段或改变数组/标量形态属于 Breaking 变化，需要新目标索引与回放迁移。",
        "上线顺序为安装不可变模板和 Pipeline、模拟摄取、样例回放、影子运行、比较结果、切换写入或读取别名，再停用旧版本。",
    ])

    doc.add_heading("9 后端实现验收", level=1)
    add_bullets(doc, [
        "所有 index template 均能被 Elasticsearch 接受，组合后的 Mapping 无冲突。",
        "入口事件被路由到唯一领域；非法域、缺关键字段或不支持来源进入 quarantine。",
        "event.id 与派生对象 ID 支持幂等回放，不因重复摄取增加行为计数。",
        "用户级网络异常只有在事件时刻存在足够置信度的用户 设备 IP 关系时才允许产生。",
        "Mapping、解析、关系解析、特征、模型和规则版本可从异常与风险结果追溯。",
        "ILM 删除时间不短于基线训练窗口和调查证据回溯要求。",
    ])
    return doc


def doc_mapping():
    doc = Document()
    configure(doc, "UEBA 场景要素与 Index 字段对应关系")
    add_title(doc, "UEBA 场景要素与 Index 字段对应关系", "首期检测用例的数据合同和降级规则")
    add_para(doc, "本文把首期 UEBA 场景拆解为可执行的数据合同。每个场景必须明确输入事件域、关键字段、实体关系、特征、基线、异常和风险输出。字段出现不等于场景可用；只有事件质量、字段覆盖、关系置信度和版本条件同时满足，场景才可标记为 Ready。")

    doc.add_heading("1 对应关系模型", level=1)
    add_code(doc, "场景\n  → 输入事件域与 ueba.event.type\n  → 必需事实字段与质量 usable_for\n  → 实体及事件时态关系\n  → feature.id 与窗口\n  → baseline.id 与 model.version\n  → anomaly.type 与证据引用\n  → risk event 与当前风险\n  → case 与分析反馈")
    add_para(doc, "检测任务不能只按索引名或 event.action 判断可用性。建议为每个场景保存一份可版本化合同，至少包含 required_event_types、required_fields、required_relations、quality_gate、feature_definitions、baseline_policy、rule_version 和 output_entity_type。")

    doc.add_heading("2 场景与索引总览", level=1)
    overview = [
        ("异常登录时间", "authentication", "account person", "features baselines anomalies risk"),
        ("登录失败量异常", "authentication", "account person", "features baselines anomalies risk"),
        ("非常用登录来源", "authentication session", "account IP person", "relations features baselines anomalies risk"),
        ("非常用设备登录", "authentication session", "account device person", "entities relations features baselines anomalies risk"),
        ("特权账号异常登录", "authentication iam", "privileged account person", "entities features baselines anomalies risk"),
        ("新外部目的地", "network", "device IP 可选 person", "relations features baselines anomalies risk"),
        ("异常外发字节量", "network", "device IP 可选 person", "relations features baselines anomalies risk"),
        ("稀有 DNS 域名", "dns", "device IP 可选 person", "relations features baselines anomalies risk"),
        ("罕见 HTTP User Agent 或 URL", "web", "device IP 可选 person", "relations features baselines anomalies risk"),
        ("新进程或罕见进程", "endpoint", "device account 可选 person", "entities features baselines anomalies risk"),
        ("敏感组成员变化", "iam directory", "actor target group person", "entities features anomalies risk cases"),
        ("疑似分片上传", "web network file", "device IP 可选 person", "upload sessions features baselines anomalies risk"),
    ]
    add_table(doc, ["场景", "输入事件域", "输出主体", "涉及派生索引"], overview, [1.55, 1.3, 1.45, 2.3], 8.0)

    doc.add_heading("3 场景字段合同", level=1)
    scenarios = [
        {
            "name":"3.1 异常登录时间",
            "events":"logs-ueba.authentication-*；ueba.event.type=authentication.login；event.outcome=success",
            "required":"@timestamp、event.id、organization.id、user.id 或可规范化的 user.name/domain",
            "conditional":"source.ip、host.id/hostname、ueba.time.quality；远程登录要求来源地址",
            "relation":"账号到人员关系用于人员输出；仅账号评分时可不解析 person",
            "feature":"auth.login.hour_histogram、auth.login.off_hours_count；窗口按账号和租户分组",
            "baseline":"账号或同伴组的登录时段分布，保存训练区间、样本量和 model.version",
            "output":"anomaly.type=unusual_login_time；evidence.event_ids 引用登录事件；风险写入 account，关系可靠时可汇总 person",
            "degrade":"时间使用 fallback、账号缺失或样本不足时 unavailable；不得用接收时间伪装行为时间",
        },
        {
            "name":"3.2 登录失败量异常",
            "events":"logs-ueba.authentication-*；authentication.login；event.outcome=failure",
            "required":"@timestamp、event.id、event.outcome、user 标识；统计源可为 Windows 或 VPN",
            "conditional":"source.ip、host.id、event.code、event.reason 用于拆分原因和调查",
            "relation":"账号到人员关系用于人员风险；来源 IP 可作为独立实体",
            "feature":"auth.failure_count、auth.distinct_source_count、auth.failure_then_success；固定 5m 1h 24h 窗口",
            "baseline":"账号自身与同伴组失败量分布；记录输入事件数和缺失率",
            "output":"anomaly.type=login_failure_volume 或 failure_then_success；风险原因引用规则与窗口",
            "degrade":"结果字段缺失、重复控制失败或采集窗口不完整时禁止产出 未发现异常",
        },
        {
            "name":"3.3 非常用登录来源",
            "events":"authentication 或 session；成功登录",
            "required":"@timestamp、user 标识、source.ip、event.outcome、event.id",
            "conditional":"source.geo、network.type、ueba.session.type、VPN 会话信息用于解释",
            "relation":"账号 person 必要于人员输出；VPN 分配地址关系应按 valid_from valid_to 解析",
            "feature":"auth.source_first_seen、auth.source_frequency、auth.distinct_sources",
            "baseline":"账号历史来源集合或网段 地理 ASN 分布，明确冷启动策略",
            "output":"anomaly.type=unusual_login_source；evidence 同时引用认证和可用的会话关系",
            "degrade":"source.ip 缺失则场景 unavailable；不能把 NAT 出口长期等同于个人设备",
        },
        {
            "name":"3.4 非常用设备登录",
            "events":"authentication；可辅以 session 与资产状态",
            "required":"@timestamp、user 标识、host.id/hostname 或等价稳定设备标识、event.outcome=success",
            "conditional":"source.ip、host.ip、设备共享标志、资产 criticality",
            "relation":"account uses device 和 account belongs_to person；关系必须覆盖事件时刻",
            "feature":"auth.device_first_seen、auth.device_frequency、auth.distinct_devices",
            "baseline":"账号常用设备集合；共享设备与跳板机需单独策略",
            "output":"anomaly.type=new_or_rare_device_login；风险先归账号/设备，满足关系门槛后归人员",
            "degrade":"只有 IP 而没有稳定设备键时最多输出 非常用来源，不能声称新设备",
        },
        {
            "name":"3.5 特权账号异常登录",
            "events":"authentication + iam；Windows 4672 或身份主档提供特权属性",
            "required":"@timestamp、account 实体、登录结果、特权标识或角色、event.id",
            "conditional":"source.ip、host.id、登录类型、session.id、目标系统 criticality",
            "relation":"account belongs_to person；privileged_role assigned_to account；目标设备关系",
            "feature":"priv.login_count、priv.off_hours_count、priv.new_source、priv.new_device",
            "baseline":"特权账号单独基线，禁止与普通账号同池训练",
            "output":"anomaly.type=privileged_account_login_anomaly；risk reason_codes 保留特权上下文",
            "degrade":"无法证明账号特权属性时不能输出特权场景，只能进入普通登录检测",
        },
        {
            "name":"3.6 新外部目的地",
            "events":"logs-ueba.network-*；网络连接建立或可判定的通信事实",
            "required":"@timestamp、source.ip、destination.ip、event.id、network.direction 或可验证的方向规则",
            "conditional":"destination.port、network.protocol、community_id、observer.*、字节包数",
            "relation":"source IP assigned_ip device；用户输出还需 device logged_on account/person 的时态链",
            "feature":"net.destination_first_seen、net.destination_frequency、net.distinct_external_destinations",
            "baseline":"设备或用户历史目的地集合；内外网边界和允许列表必须版本化",
            "output":"anomaly.type=new_external_destination；设备异常始终可独立存在，人员风险需关系门禁",
            "degrade":"方向未知或 NAT/共享地址无法归属时只做 IP/设备级分析",
        },
        {
            "name":"3.7 异常外发字节量",
            "events":"network；优先 Zeek conn 或能提供方向字节的等价来源",
            "required":"@timestamp、source.ip、destination.ip、source.bytes 或可靠方向化 network.bytes、event.id",
            "conditional":"destination.port、protocol、duration、packets、observer、质量告警",
            "relation":"同新外部目的地；必须使用事件时刻的设备 IP 和用户设备关系",
            "feature":"net.outbound_bytes_sum、net.outbound_bytes_by_destination、net.upload_ratio",
            "baseline":"按实体、时段和业务日历的外发量分布；保存单位 bytes",
            "output":"anomaly.type=abnormal_outbound_volume；explanation 保存当前值、期望值和偏离",
            "degrade":"Zeek missed_bytes 或方向不可信时 partial 且禁止依赖精确流量的规则",
        },
        {
            "name":"3.8 稀有 DNS 域名",
            "events":"logs-ueba.dns-*；DNS query/response",
            "required":"@timestamp、source.ip、dns.question.name、event.id",
            "conditional":"dns.question.type、response_code、resolved_ip、observer、destination.ip",
            "relation":"source IP 到设备；用户输出需完整时态归属链",
            "feature":"dns.domain_first_seen、dns.domain_frequency、dns.distinct_domains、dns.nxdomain_rate",
            "baseline":"设备 用户 同伴组域名频率；域名规范化与公共后缀规则要版本化",
            "output":"anomaly.type=rare_dns_domain；证据引用 DNS 事件，可关联后续网络连接",
            "degrade":"查询名缺失时 unavailable；仅有解析器汇总指标不能替代逐查询行为",
        },
        {
            "name":"3.9 罕见 HTTP User Agent 或 URL",
            "events":"logs-ueba.web-*；HTTP 请求",
            "required":"@timestamp、source.ip、event.id，以及 user_agent.original 或 url.original/full 至少一项",
            "conditional":"http.request.method、status_code、url.domain/path/query、destination、字节数",
            "relation":"source IP 到设备；用户级输出需 account/person 的事件时态关系",
            "feature":"web.user_agent_frequency、web.url_frequency、web.new_domain、web.request_count",
            "baseline":"对 UA 和 URL 分开建模；URL 参数需脱敏和规范化后再做基线",
            "output":"anomaly.type=rare_user_agent 或 rare_url；证据引用 Web 事件",
            "degrade":"TLS 未解密且无代理日志时不可宣称 URL 场景可用；只有域名时降级为域名场景",
        },
        {
            "name":"3.10 新进程或罕见进程",
            "events":"logs-ueba.endpoint-*；Windows 4688 或 EDR process start",
            "required":"@timestamp、host.id/hostname、process.executable 或 process.name、event.id",
            "conditional":"user、parent process、command_line、args、hash、签名和目标设备 criticality",
            "relation":"host 对应 device；user/account 到 person 用于人员输出",
            "feature":"proc.first_seen、proc.frequency、proc.parent_child_frequency、proc.command_pattern_frequency",
            "baseline":"按设备、设备类型、用户和同伴组建立；路径大小写和参数需规范化",
            "output":"anomaly.type=new_or_rare_process；可分别归设备和账号，证据保留命令行引用",
            "degrade":"缺进程名/路径则 unavailable；缺父进程或命令行时仅降级相关子规则",
        },
        {
            "name":"3.11 敏感组成员变化",
            "events":"logs-ueba.iam-* 或 directory-*；group membership change",
            "required":"@timestamp、actor user、target user/account、group.id/name、event.action/type、event.id",
            "conditional":"host、source.ip、event.outcome、reason、敏感组标签和审批引用",
            "relation":"actor performs change；target member_of group；account belongs_to person",
            "feature":"iam.sensitive_group_change_count、iam.actor_target_pair_frequency；可直接规则检测",
            "baseline":"管理员正常变更范围与时段；敏感组目录由治理配置版本化",
            "output":"anomaly.type=sensitive_group_membership_change；风险可同时归 actor、target 和 group，并创建 case",
            "degrade":"不能区分 Subject 与 Target 时不得归责；缺敏感组目录时只能记录普通组变更",
        },
        {
            "name":"3.12 疑似分片上传",
            "events":"logs-ueba.web-* 为主，关联 network-* 与 file-*；web.request 可附 upload 语义标签",
            "required":"@timestamp、event.id、organization.id、source.ip 或设备键、destination.ip/domain、http.request.method、http.request.body.bytes 或可靠出站字节",
            "conditional":"upload session ID、part number/count、Content Range、URL、file ID/name/size/hash、Zeek orig_fuids、status code",
            "relation":"source IP assigned_ip device；用户输出还需事件时刻有效的 device logged_on account/person 关系",
            "feature":"upload.part_count、upload.bytes_sum、upload.concurrent_sessions、upload.new_destination、upload.retry_ratio；窗口按会话与实体分组",
            "baseline":"设备或用户的上传量、分片数、目的地和时间分布；应用层正文与网络层字节分别建模",
            "output":"有稳定 upload.session.id 时写入 ueba-upload-sessions-* 并聚合完整度；异常输出 chunked_upload_suspected 或 abnormal_chunked_upload，证据引用所有相关分片",
            "degrade":"缺会话 ID 时仅做时间窗口启发式关联并标记 partial/低置信；无法证明文件 ID、总大小或哈希时不得声称完整文件已上传或具体文件已泄露",
        },
    ]
    for s in scenarios:
        doc.add_heading(s["name"], level=2)
        rows = [
            ("输入事件", s["events"]),
            ("必需字段", s["required"]),
            ("条件字段", s["conditional"]),
            ("实体关系", s["relation"]),
            ("特征索引", s["feature"]),
            ("基线索引", s["baseline"]),
            ("异常和风险", s["output"]),
            ("降级门禁", s["degrade"]),
        ]
        add_table(doc, ["要素", "对应关系"], rows, [1.15, 5.45], 8.4)
        if s["name"].startswith("3.12"):
            add_para(doc, "首期能力边界如下。分片字段尚未进入正式公共 Mapping 时，可先保存在 vendor.payload；输出只能标记为疑似分片上传或异常外发，不得绕过 strict Mapping，也不得把启发式聚合解释为确定的完整文件上传。")
            capability = [
                ("单个上传请求识别", "可支持", "方法、正文长度或文件关联足以证明上传"),
                ("分片请求字节统计", "可支持", "明确应用层或网络层口径，不重复相加"),
                ("短时间大量分片上传", "有条件支持", "窗口完整且主体与目的地可识别"),
                ("多个分片归并为一次上传", "有条件支持", "必须获得稳定 upload.session.id 或等价厂商字段"),
                ("完整文件大小和哈希", "暂不保证", "只有来源提供完整文件事实时才能填写"),
                ("重传和断点续传去重", "暂不保证", "需要 part number、range、ETag 或等价幂等键"),
                ("具体文件外泄判定", "不支持直接定性", "首期只能输出疑似上传或异常外发，需要 DLP、文件审计等证据增强"),
            ]
            add_table(doc, ["能力", "首期状态", "数据条件"], capability, [2.0, 1.15, 3.45], 8.3)

    doc.add_heading("4 跨场景公共字段", level=1)
    shared = [
        ("租户", "organization.id", "所有查询、聚合、关系和风险必须按租户隔离"),
        ("事件主键", "event.id", "幂等摄取、证据引用、回放去重"),
        ("稳定语义", "ueba.event.type", "检测优先依赖受控事件类型，不依赖自由字符串 event.action"),
        ("可用性", "ueba.quality.status usable_for warnings", "按场景与子能力门禁，不把 partial 一刀切"),
        ("时间可信度", "ueba.time.quality source", "顺序和窗口敏感场景必须拒绝低质量时间"),
        ("处理版本", "ueba.provenance.mapping_version parser_version", "回放、对比和结果解释"),
        ("证据位置", "raw_event_id raw_uri raw_sha256", "调查时回查不可变原始证据"),
        ("实体版本", "relation.resolution_version inputs.resolution_version", "关系修正后定位需重算结果"),
        ("模型与规则", "model.version detection.rule_version calculation.policy_version", "保证检测结果可复现"),
    ]
    add_table(doc, ["合同要素", "字段", "约束"], shared, [1.1, 2.65, 2.85], 8.2)

    doc.add_heading("5 数据就绪状态", level=1)
    readiness = [
        ("Ready", "事件域存在；必需字段覆盖达标；时间可信；关系置信度达标；窗口完整", "正常运行并输出指定实体类型"),
        ("Degraded", "条件字段或部分子能力缺失，但核心事实仍可计算", "运行受支持的子规则，附 quality warning"),
        ("Unavailable", "必需字段、事件源、时间、关系或采集窗口不满足", "停止场景；输出数据不可用，不能输出未发现异常"),
        ("Unsupported", "来源或事件类型未进入映射目录", "进入 quarantine 或仅保留原始证据"),
    ]
    add_table(doc, ["状态", "判定", "检测行为"], readiness, [1.05, 3.65, 1.9], 8.4)

    doc.add_heading("6 建议的场景合同记录", level=1)
    add_code(doc, "use_case.id / version / owner\nrequired_event_types[]\nrequired_fields[] / conditional_fields[]\nrequired_relations[] + min_confidence\nquality.required_usable_for[]\nquality.max_time_fallback_rate\nfeature.id / version / window / dimensions\nbaseline.model_id / model_version / min_samples\ndetection.rule_id / rule_version\noutput.entity_types[] / anomaly_type / risk_policy\ndegradation_policy / test_cases[]")
    add_para(doc, "这份合同建议与 Mapping、路由注册表和测试样例一起发布。任何必需字段、关系门槛、特征定义或输出主体的变化，都应触发场景版本升级和历史回放验证。")

    doc.add_heading("7 验收清单", level=1)
    add_bullets(doc, [
        "每个场景至少有正常、异常、缺字段、未知枚举、重复事件、乱序和来源版本变化样例。",
        "事件字段能够定位到具体组件模板，场景输出能够定位到 feature、baseline、anomaly 和 risk 字段。",
        "关系 valid_from valid_to 覆盖事件时间，过期或低置信度关系不会把设备异常错误归属到人员。",
        "检测输入为零或 capability 不可用时返回数据不可用，而不是未发现异常。",
        "重复回放不增加独立动作数；映射或关系版本变化后可定位并重算受影响窗口。",
        "异常解释包含当前值、期望值、偏离、原因码、证据 ID、模型和规则版本。",
        "分片上传测试覆盖正常多片、缺片、乱序、重复分片、断点续传、并发上传、缺 session ID 和方向不明等情况。",
    ])
    return doc


if __name__ == "__main__":
    d1 = doc_structure()
    d1.save(OUT / "UEBA后端Index表结构设计.docx")
    d2 = doc_mapping()
    d2.save(OUT / "UEBA场景要素与Index字段对应关系.docx")
    print(OUT / "UEBA后端Index表结构设计.docx")
    print(OUT / "UEBA场景要素与Index字段对应关系.docx")
